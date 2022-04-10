from tkinter import *
import tkinter.messagebox
import xlsxwriter
from labels import spacing,create
from docx import Document

#read criterias from text file
with open("criteria.txt","r") as f:
    criteria_list = [i.strip().split("-") for i in f.readlines() if i]
print(criteria_list)

if criteria_list[-1]== ['']:
    criteria_list.pop(-1)

criteria = [". ".join(i) for i in criteria_list]

print(criteria)
all_content = []
window = Tk()
window.geometry("600x400")

#spacings
spacing(0,0)
spacing(0,2)

#create header labels
create("Process",1,1)
create("Conformance",1,3)
create("Section A",0,1)
process_list = []

class Process:
    def __init__(self,text,row,process,conformance):
        self.text = text
        self.row = row
        self.process = process
        self.conformance = conformance
        self.v = 0
        self.content = [self.process,self.conformance]


    def set_description(self):
        Label(text=self.text).grid(row=self.row,column=1)
        self.v = IntVar()

        #new window popup if not conformed
        def not_conformed():

            self.content = [self.process,self.conformance,"NO"]
            new_window = Toplevel(window)
            new_window.geometry("600x400")

            Label(new_window,text=self.conformance).place(x=40,y=40)
            Label(new_window,text="Findings Description").place(x=40,y=120)
            #textbox for description of findings
            description_textbox = Text(new_window,width=24,height=8)
            description_textbox.place(x=340,y=100)

            #classification labels
            Label(new_window,text="Classification of finding").place(x=40,y=260)
            Label(new_window,text="Please select one ").place(x=340,y=260)

            #dropdown selection menu
            opt_btn = StringVar()
            options = OptionMenu(new_window, opt_btn, "     MAJOR     ", "     MINOR     ", "OBSERVATION")
            options.place(x=450, y=260)

            def description():
                if description_textbox.get("1.0","end") and opt_btn.get():
                    self.content.append(description_textbox.get("1.0","end"))
                    self.content.append(opt_btn.get())
                else:
                    tkinter.messagebox.showwarning(title="Warning",message="Please fill in the necessary information")
                    return
                new_window.destroy()

            Button(new_window,text="Save",command=description).place(x=340,y=40)
            new_window.mainloop()

        def conformed():
            self.content = [self.process,self.conformance,"YES"]




        Radiobutton(window,text="NO",variable=self.v,value=2,command=not_conformed).grid(row=self.row,column=4)
        Radiobutton(window,text="YES",variable=self.v,value=1,command=conformed).grid(row=self.row,column=3)

for i in range(len(criteria)):
    process_list.append(Process(criteria[i],2*i+1,criteria_list[i][0],criteria_list[i][1]))
    process_list[-1].set_description()
    spacing(2*i+2,0)

def add():
    new_window = Toplevel(window)
    new_window.geometry("600x400")
    Label(new_window,text="Section & Criteria Number").place(x=40,y=40)
    Label(new_window,text="Criteria").place(x=40,y=120)

    section_text = Text(new_window, width=10, height=1)
    section_text.place(x=200, y=40)
    criteria_text = Text(new_window, width=20, height=2)
    criteria_text.place(x=200, y=120)

    def save():
        section = section_text.get(1.0, END)
        new_criteria = criteria_text.get(1.0, END)
        with open("criteria.txt", "a") as file:
            if section and new_criteria:
                file.write(f"{section}-{new_criteria}\n")
            else:
                tkinter.messagebox.showwarning(title="Warning", message="Please fill in the necessary information")
        new_window.destroy()

    Button(new_window,text="Save",command=save).place(x=200,y=300)

def save():
    all_content.append([])
    for item in process_list:
        if item.content[2] in ["YES","NO"]:
            all_content[-1].append(item.content)
        else:
            all_content.pop(-1)
            tkinter.messagebox.showwarning(title="Warning",message="Please fill in the necessary information")
            return
    for item in process_list:
        item.v.set(0)
        item.content = item.content[0:2]+[' ']

def export():
    if all_content:
        pass
    else:
        tkinter.messagebox.showwarning(title="Warning",message="Please fill in the necessary information")

    #export to excel sheet
    wk=xlsxwriter.Workbook("checklist.xlsx")
    ws=wk.add_worksheet("sheet")
    j=1
    for item in all_content:
        for k in item:
            ws.write_row(j,0,k)
            j+=1

    wk.close()

    #export to word doc

    document = Document()
    #create new header
    table = document.add_table(rows=1,cols=5,style="Table Grid")
    for item in all_content:
        for k in item:
            row_cells = table.add_row().cells
            for j in range(len(k)):
                row_cells[j].text = k[j]

    document.save("checklist.docx")

Button(window,text="Save",command=save).grid(row=2*len(criteria),column=5)

Button(window,text="Export",command=export).grid(row=3*len(criteria)+1,column=5)

Button(window,text="Add Criteria",command=add).grid(row=2*len(criteria),column=1)







window.mainloop()